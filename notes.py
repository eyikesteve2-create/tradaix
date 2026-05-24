import streamlit as st
import pandas as pd
import json
import math
import re
import io
import os
import requests
from datetime import datetime
from dotenv import load_dotenv
from db import get_db
from inspinote import lister_modeles, suggerer_modele, construire_contexte_modele

load_dotenv()

# ─────────────────────────────────────────────
# CONSTANTES
# ─────────────────────────────────────────────
CATEGORIES = ["Général", "Réunion", "Finance", "Projet", "RH", "Commercial", "Technique"]
COULEURS   = {
    "Général":    "#1E3A8A", "Réunion":  "#7C3AED", "Finance":  "#059669",
    "Projet":     "#CC0000", "RH":       "#F59E0B",  "Commercial":"#0EA5E9",
    "Technique":  "#64748B",
}

TYPES_DOCUMENT = {
    "📝 Note":           "note interne d'entreprise",
    "📋 Compte-rendu":   "compte-rendu de réunion professionnel",
    "📊 Rapport":        "rapport d'activité ou d'analyse détaillé",
    "📌 Mémo":           "mémo interne court et direct",
    "📈 Bilan":          "bilan de performance ou de résultats",
    "📁 Synthèse":       "synthèse structurée d'informations",
}

# ─────────────────────────────────────────────
# CSS
# ─────────────────────────────────────────────
def inject_css():
    st.markdown("""
    <style>
    .note-card {
        background:#fff; border-radius:12px; padding:14px 18px; margin-bottom:10px;
        border-left:4px solid #1E3A8A; box-shadow:0 2px 8px rgba(0,0,0,0.07);
    }
    .note-titre  { font-weight:700; color:#1E293B; font-size:15px; margin-bottom:4px; }
    .note-meta   { font-size:11px; color:#888; margin-bottom:6px; }
    .note-apercu { font-size:13px; color:#555; white-space:pre-wrap; max-height:60px; overflow:hidden; }
    .cat-badge   { display:inline-block; padding:2px 10px; border-radius:10px;
                   font-size:11px; font-weight:600; color:#fff; margin-right:6px; }
    .note-editor { background:#FAFBFC; border-radius:10px; padding:16px;
                   border:1px solid #E5E7EB; margin-top:8px; }
    .ia-box      { background:#F0F7FF; border-radius:10px; padding:16px;
                   border:1px solid #BFDBFE; margin-bottom:16px; }
    .calc-result { background:#F0FDF4; border-radius:8px; padding:10px 14px;
                   font-size:14px; color:#065F46; font-weight:600; margin-top:8px; }
    .calc-error  { background:#FEF2F2; border-radius:8px; padding:10px 14px;
                   font-size:14px; color:#CC0000; margin-top:8px; }
    .table-info  { background:#EFF6FF; border-radius:8px; padding:8px 12px;
                   font-size:12px; color:#1E40AF; margin-bottom:8px; }
    .stat-note   { background:#fff; border-radius:12px; padding:16px 10px; text-align:center;
                   box-shadow:0 2px 10px rgba(0,0,0,0.07); border-bottom:3px solid #E5E7EB; }
    .stat-note-num   { font-size:2em; font-weight:800; }
    .stat-note-label { font-size:12px; color:#666; margin-top:4px; }
    </style>
    """, unsafe_allow_html=True)


# ─────────────────────────────────────────────
# GÉNÉRATION IA AVEC MODÈLES INSPINOTE
# ─────────────────────────────────────────────
def generer_note_ia(type_doc: str, sujet: str, categorie: str,
                    contexte: str = "", modele_force: dict = None) -> dict:
    """Génère titre + contenu + tableau via Groq en s'inspirant d'un modèle inspinote."""
    groq_key   = os.getenv("GROQ_API_KEY")
    type_label = TYPES_DOCUMENT.get(type_doc, "document professionnel")

    # ── Sélection du modèle ──────────────────
    if modele_force:
        modele = modele_force
    else:
        modele = suggerer_modele(type_doc, sujet)

    contexte_modele = construire_contexte_modele(modele) if modele else ""
    info_modele     = f"(modèle : {modele['nom']})" if modele else "(sans modèle)"

    prompt = f"""Tu es un assistant rédacteur professionnel pour l'entreprise TradAIx.
Rédige un {type_label} complet sur le sujet suivant : "{sujet}"
Catégorie : {categorie}
{f'Contexte supplémentaire : {contexte}' if contexte else ''}

{contexte_modele}

INSTRUCTIONS :
- Respecte ABSOLUMENT la structure et les sections du modèle de référence ci-dessus
- Remplace les balises [ENTRE CROCHETS] par le contenu réel adapté au sujet
- Rédige en français professionnel
- Génère un tableau pertinent avec de vraies données adaptées au sujet

Réponds UNIQUEMENT avec un JSON valide (sans markdown) :
{{
  "titre": "titre du document (court et précis)",
  "contenu": "contenu complet structuré avec sections claires séparées par sauts de ligne",
  "tableau_entetes": ["Colonne1", "Colonne2", "Colonne3"],
  "tableau_donnees": [["val1","val2","val3"],["val1","val2","val3"]]
}}

Si aucun tableau n'est pertinent, mets tableau_entetes=[] et tableau_donnees=[]."""

    try:
        r = requests.post(
            "https://api.groq.com/openai/v1/chat/completions",
            headers={"Authorization": f"Bearer {groq_key}", "Content-Type": "application/json"},
            json={
                "model": "llama-3.3-70b-versatile",
                "messages": [{"role": "user", "content": prompt}],
                "max_tokens": 2500,
                "temperature": 0.4,
            },
            timeout=30
        )
        contenu_brut = r.json()["choices"][0]["message"]["content"].strip()
        if "```" in contenu_brut:
            contenu_brut = contenu_brut.split("```")[1]
            if contenu_brut.startswith("json"):
                contenu_brut = contenu_brut[4:]
        data    = json.loads(contenu_brut.strip())
        entetes = data.get("tableau_entetes", [])
        donnees = data.get("tableau_donnees", [])
        tableau = [entetes] + donnees if entetes else []
        return {
            "titre":        data.get("titre", sujet),
            "contenu":      data.get("contenu", ""),
            "tableau":      tableau,
            "modele_utilise": info_modele,
        }
    except Exception as e:
        return {
            "titre":        sujet,
            "contenu":      f"Erreur lors de la génération : {e}",
            "tableau":      [],
            "modele_utilise": info_modele,
        }


# ─────────────────────────────────────────────
# DB — CRUD NOTES
# ─────────────────────────────────────────────
def get_notes(user_email: str) -> list:
    res = get_db().table("notes").select("*").eq("cree_par", user_email)\
              .order("updated_at", desc=True).execute()
    return res.data or []

def sauvegarder_note(user_email: str, note: dict) -> int:
    db  = get_db()
    now = datetime.now().isoformat()
    data = {
        "titre":      note.get("titre", "Sans titre")[:100],
        "contenu":    note.get("contenu", ""),
        "categorie":  note.get("categorie", "Général"),
        "tableau":    json.dumps(note.get("tableau", []), ensure_ascii=False),
        "cree_par":   user_email,
        "updated_at": now,
    }
    nid = note.get("id")
    if nid:
        db.table("notes").update(data).eq("id", nid).execute()
        return nid
    else:
        data["created_at"] = now
        res = db.table("notes").insert(data).execute()
        return res.data[0]["id"] if res.data else None

def supprimer_note(nid: int):
    get_db().table("notes").delete().eq("id", nid).execute()


# ─────────────────────────────────────────────
# EXPORT WORD (.docx)
# ─────────────────────────────────────────────
def exporter_note_word(note: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3)
        section.right_margin  = Cm(2)

    # En-tête
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.text = "TradAIx — Document Professionnel"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.runs[0].font.size = Pt(9)
    hp.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    # Titre
    titre_para = doc.add_paragraph()
    run_titre = titre_para.add_run(note.get("titre", "Sans titre"))
    run_titre.bold = True
    run_titre.font.size = Pt(20)
    run_titre.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    p = titre_para._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1E3A8A")
    pBdr.append(bottom)
    pPr.append(pBdr)

    # Métadonnées
    meta = doc.add_paragraph()
    r1 = meta.add_run(f"Catégorie : {note.get('categorie','Général')}     ")
    r1.font.size = Pt(10); r1.bold = True
    r1.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    r2 = meta.add_run(f"Date : {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    doc.add_paragraph()

    # Contenu
    contenu = note.get("contenu", "").strip()
    if contenu:
        for ligne in contenu.split("\n"):
            p = doc.add_paragraph()
            if ligne.strip().startswith("##") or ligne.strip().startswith("**"):
                run = p.add_run(ligne.replace("##","").replace("**","").strip())
                run.bold = True
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            else:
                run = p.add_run(ligne)
                run.font.size = Pt(11)

    doc.add_paragraph()

    # Tableau
    tableau_json = note.get("tableau", [])
    if isinstance(tableau_json, str):
        try:    tableau_json = json.loads(tableau_json)
        except: tableau_json = []

    if tableau_json and len(tableau_json) > 1:
        h2 = doc.add_heading("Tableau de données", level=2)
        h2.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        entetes = tableau_json[0]
        donnees = tableau_json[1:]
        nb_col  = len(entetes)
        table   = doc.add_table(rows=1 + len(donnees), cols=nb_col)
        table.style     = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for j, entete in enumerate(entetes):
            cell = table.rows[0].cells[j]
            cell.text = str(entete)
            run = cell.paragraphs[0].runs[0]
            run.bold = True; run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            tc = cell._tc; tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "1E3A8A"); tcPr.append(shd)

        for i, row in enumerate(donnees):
            for j in range(nb_col):
                val  = row[j] if j < len(row) else ""
                cell = table.rows[i+1].cells[j]
                cell.text = str(val)
                if i % 2 == 0:
                    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "EFF6FF"); tcPr.append(shd)

    # Pied de page
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.text = f"TradAIx © {datetime.now().year} — Document confidentiel"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.runs[0].font.size = Pt(9)
    fp.runs[0].font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# ─────────────────────────────────────────────
# CALCUL
# ─────────────────────────────────────────────
def evaluer_expression(expr: str) -> str:
    try:
        expr_clean = expr.strip().replace(",", ".").replace("^", "**")
        autorise = {
            "__builtins__": {}, "abs": abs, "round": round, "min": min, "max": max,
            "sum": sum, "pow": pow, "sqrt": math.sqrt, "log": math.log,
            "log10": math.log10, "sin": math.sin, "cos": math.cos,
            "tan": math.tan, "pi": math.pi, "e": math.e,
        }
        result = eval(expr_clean, autorise)
        return f"{result:,.4f}".rstrip("0").rstrip(".") if isinstance(result, float) else str(result)
    except ZeroDivisionError:
        return "❌ Division par zéro"
    except Exception as e:
        return f"❌ Erreur : {e}"


# ─────────────────────────────────────────────
# FORMULAIRE GÉNÉRATION IA
# ─────────────────────────────────────────────
def formulaire_generation_ia(user_email: str):
    st.markdown('<div class="ia-box">', unsafe_allow_html=True)
    st.markdown("#### 🤖 Générer un document avec l'IA")

    col1, col2 = st.columns(2)
    with col1:
        type_doc  = st.selectbox("📄 Type de document", list(TYPES_DOCUMENT.keys()), key="ia_type")
    with col2:
        categorie = st.selectbox("🗂️ Catégorie", CATEGORIES, key="ia_cat")

    sujet    = st.text_input("✏️ Sujet du document",
                             placeholder="Ex: Réunion budgétaire Q2 2026, Bilan RH mars 2026...",
                             key="ia_sujet")
    contexte = st.text_area("📎 Contexte supplémentaire (optionnel)",
                            placeholder="Ex: Participants : Jean, Marie. Points abordés : budget, recrutement...",
                            height=80, key="ia_contexte")

    # ── SÉLECTEUR DE MODÈLE INSPINOTE ────────
    st.markdown("---")
    st.markdown("**📁 Modèle de référence (inspinote)**")
    modeles_dispo = lister_modeles()

    col_m1, col_m2 = st.columns([3, 1])
    with col_m1:
        # Suggestion automatique basée sur le type
        suggestion = suggerer_modele(type_doc, sujet or "") if modeles_dispo else None
        options_modeles = ["🤖 Auto (suggestion IA)"] + [m["nom"] for m in modeles_dispo]
        idx_defaut = 0
        if suggestion:
            try:
                idx_defaut = [m["nom"] for m in modeles_dispo].index(suggestion["nom"]) + 1
            except ValueError:
                idx_defaut = 0
        choix_modele = st.selectbox("Choisir un modèle", options_modeles,
                                    index=idx_defaut, key="ia_modele_choix")
    with col_m2:
        if suggestion and choix_modele == "🤖 Auto (suggestion IA)":
            st.info(f"💡 Suggéré :\n**{suggestion['nom']}**")
        elif modeles_dispo:
            st.success(f"📄 {len(modeles_dispo)} modèle(s)")

    # Aperçu du modèle sélectionné
    modele_selectionne = None
    if modeles_dispo:
        if choix_modele == "🤖 Auto (suggestion IA)":
            modele_selectionne = None  # sera auto-sélectionné dans generer_note_ia
        else:
            for m in modeles_dispo:
                if m["nom"] == choix_modele:
                    modele_selectionne = m
                    break
        if modele_selectionne:
            with st.expander(f"👁️ Aperçu du modèle : {modele_selectionne['nom']}"):
                st.text(modele_selectionne["contenu"][:800] + "...")

    if not modeles_dispo:
        st.warning("⚠️ Aucun modèle trouvé dans le dossier `inspinote/`. L'IA rédigera sans modèle.")

    if st.button("🚀 Générer le document", use_container_width=True, type="primary", key="ia_generer"):
        if not sujet.strip():
            st.error("Veuillez entrer un sujet.")
        else:
            with st.spinner("✍️ L'IA consulte le modèle et rédige votre document..."):
                result = generer_note_ia(type_doc, sujet, categorie, contexte,
                                         modele_force=modele_selectionne)
                st.session_state["ia_result"] = {**result, "categorie": categorie}
            modele_info = result.get("modele_utilise", "")
            st.success(f"✅ Document généré ! {modele_info}")

    st.markdown('</div>', unsafe_allow_html=True)

    # Afficher le résultat généré
    if st.session_state.get("ia_result"):
        r = st.session_state["ia_result"]
        st.markdown("---")
        st.markdown("#### 📄 Document généré — Vérifiez et modifiez si besoin")

        col1, col2 = st.columns([3, 1])
        with col1:
            titre = st.text_input("Titre", value=r.get("titre",""), key="ia_res_titre")
        with col2:
            cat = st.selectbox("Catégorie", CATEGORIES,
                               index=CATEGORIES.index(r.get("categorie","Général")),
                               key="ia_res_cat")

        contenu = st.text_area("Contenu", value=r.get("contenu",""), height=300, key="ia_res_contenu")

        # Tableau généré
        tableau = r.get("tableau", [])
        if tableau and len(tableau) > 1:
            st.markdown("**📊 Tableau généré :**")
            try:
                df = pd.DataFrame(tableau[1:], columns=tableau[0])
                st.dataframe(df, use_container_width=True, hide_index=True)
            except Exception:
                pass

        # Calculs sur tableau
        if tableau and len(tableau) > 1:
            st.markdown("#### 🧮 Calculs sur le tableau")
            col_c1, col_c2, col_c3 = st.columns([2, 2, 1])
            with col_c1:
                col_calcul = st.selectbox("Colonne", tableau[0], key="ia_col_calc")
            with col_c2:
                operation  = st.selectbox("Opération", ["Somme","Moyenne","Maximum","Minimum","Compte"], key="ia_op")
            with col_c3:
                if st.button("▶️", key="ia_calc_btn", use_container_width=True):
                    try:
                        idx = tableau[0].index(col_calcul)
                        valeurs = [float(row[idx].replace(",",".")) for row in tableau[1:]
                                   if idx < len(row) and row[idx].strip()]
                        ops = {"Somme": sum(valeurs), "Moyenne": sum(valeurs)/len(valeurs),
                               "Maximum": max(valeurs), "Minimum": min(valeurs), "Compte": len(valeurs)}
                        st.markdown(f'<div class="calc-result">📊 {operation} de «{col_calcul}» = <strong>{ops[operation]:,.2f}</strong></div>',
                                    unsafe_allow_html=True)
                    except Exception as e:
                        st.markdown(f'<div class="calc-error">⚠️ {e}</div>', unsafe_allow_html=True)

        # Calculatrice libre
        st.markdown("#### 🧮 Calculatrice")
        col_e, col_b = st.columns([4, 1])
        with col_e:
            expr = st.text_input("", placeholder="Ex: 1500 * 12 / 100", key="ia_expr", label_visibility="collapsed")
        with col_b:
            if st.button("=", key="ia_calc_libre", use_container_width=True) and expr:
                res = evaluer_expression(expr)
                cls = "calc-error" if res.startswith("❌") else "calc-result"
                st.markdown(f'<div class="{cls}">{res}</div>', unsafe_allow_html=True)

        st.markdown("")

        # Boutons
        col_s, col_w, col_ann = st.columns([2, 2, 1])
        with col_s:
            if st.button("💾 Sauvegarder", use_container_width=True, type="primary", key="ia_save"):
                if not titre.strip():
                    st.error("Titre obligatoire.")
                else:
                    sauvegarder_note(user_email, {
                        "titre": titre, "contenu": contenu,
                        "categorie": cat, "tableau": tableau,
                    })
                    st.success("✅ Document sauvegardé !")
                    st.session_state.pop("ia_result", None)
                    st.rerun()
        with col_w:
            try:
                docx_bytes  = exporter_note_word({"titre": titre, "contenu": contenu,
                                                   "categorie": cat, "tableau": tableau})
                nom_fichier = re.sub(r"[^\w\- ]", "", titre or "document").strip().replace(" ", "_") + ".docx"
                st.download_button("📄 Télécharger Word", data=docx_bytes, file_name=nom_fichier,
                                   mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                   key="ia_word", use_container_width=True)
            except Exception as e:
                st.error(f"Erreur export : {e}")
        with col_ann:
            if st.button("✖️", key="ia_annuler", use_container_width=True):
                st.session_state.pop("ia_result", None)
                st.rerun()


# ─────────────────────────────────────────────
# CARTE NOTE
# ─────────────────────────────────────────────
def carte_note(note: dict) -> bool:
    couleur = COULEURS.get(note.get("categorie","Général"), "#1E3A8A")
    apercu  = note.get("contenu","")[:120].replace("\n"," ")
    date_m  = note.get("updated_at","")[:10]
    tableau = note.get("tableau","[]")
    has_tab = False
    try:
        t = json.loads(tableau) if isinstance(tableau, str) else tableau
        has_tab = isinstance(t, list) and len(t) > 1
    except Exception:
        pass

    st.markdown(f"""
    <div class="note-card" style="border-left-color:{couleur}">
        <div class="note-titre">📄 {note.get('titre','Sans titre')}</div>
        <div class="note-meta">
            <span class="cat-badge" style="background:{couleur}">{note.get('categorie','Général')}</span>
            {'📊' if has_tab else ''}
            &nbsp;·&nbsp; {date_m}
        </div>
        <div class="note-apercu">{apercu}{'...' if len(note.get('contenu',''))>120 else ''}</div>
    </div>""", unsafe_allow_html=True)

    col_o, col_w = st.columns([2, 1])
    with col_o:
        ouvrir = st.button("✏️ Ouvrir", key=f"open_{note['id']}", use_container_width=True)
    with col_w:
        try:
            docx_bytes  = exporter_note_word(note)
            nom_fichier = re.sub(r"[^\w\- ]", "", note.get("titre","note")).strip().replace(" ","_") + ".docx"
            st.download_button("📄 Word", data=docx_bytes, file_name=nom_fichier,
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               key=f"dl_{note['id']}", use_container_width=True)
        except Exception:
            pass
    return ouvrir


# ─────────────────────────────────────────────
# ÉDITEUR (modification d'une note existante)
# ─────────────────────────────────────────────
def editeur_note(note: dict, user_email: str):
    prefix = f"edit_{note.get('id','')}"
    st.markdown('<div class="note-editor">', unsafe_allow_html=True)

    col1, col2 = st.columns([3, 1])
    with col1:
        titre = st.text_input("📝 Titre", value=note.get("titre",""), key=f"{prefix}_titre")
    with col2:
        cat_idx = CATEGORIES.index(note.get("categorie","Général")) if note.get("categorie") in CATEGORIES else 0
        cat = st.selectbox("🗂️", CATEGORIES, index=cat_idx, key=f"{prefix}_cat")

    contenu = st.text_area("✏️ Contenu", value=note.get("contenu",""), height=200, key=f"{prefix}_contenu")

    # Tableau
    tableau_json = note.get("tableau", [])
    if isinstance(tableau_json, str):
        try:    tableau_json = json.loads(tableau_json)
        except: tableau_json = []

    if tableau_json and len(tableau_json) > 1:
        st.markdown("**📊 Tableau :**")
        try:
            df = pd.DataFrame(tableau_json[1:], columns=tableau_json[0])
            st.dataframe(df, use_container_width=True, hide_index=True)
        except Exception:
            pass

    st.markdown('</div>', unsafe_allow_html=True)

    col_s, col_w, col_d, col_a = st.columns([2, 2, 1, 1])
    with col_s:
        if st.button("💾 Sauvegarder", key=f"{prefix}_save", use_container_width=True, type="primary"):
            sauvegarder_note(user_email, {"id": note.get("id"), "titre": titre,
                                          "contenu": contenu, "categorie": cat,
                                          "tableau": tableau_json})
            st.success("✅ Sauvegardé !")
            st.session_state.pop("note_active", None)
            st.rerun()
    with col_w:
        try:
            docx_bytes  = exporter_note_word({"titre": titre, "contenu": contenu,
                                               "categorie": cat, "tableau": tableau_json})
            nom_fichier = re.sub(r"[^\w\- ]", "", titre or "note").strip().replace(" ","_") + ".docx"
            st.download_button("📄 Télécharger Word", data=docx_bytes, file_name=nom_fichier,
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               key=f"{prefix}_word", use_container_width=True)
        except Exception as e:
            st.error(f"Erreur : {e}")
    with col_d:
        if st.button("🗑️", key=f"{prefix}_del", use_container_width=True):
            supprimer_note(note["id"])
            st.session_state.pop("note_active", None)
            st.rerun()
    with col_a:
        if st.button("✖️", key=f"{prefix}_cancel", use_container_width=True):
            st.session_state.pop("note_active", None)
            st.rerun()


# ─────────────────────────────────────────────
# GESTIONNAIRE DE MODÈLES INSPINOTE
# ─────────────────────────────────────────────
def gestionnaire_modeles():
    """Permet d'uploader, voir et supprimer les modèles dans inspinote/."""
    from inspinote import DOSSIER_MODELES
    import shutil

    with st.expander("📁 Gérer les modèles inspinote", expanded=False):
        st.markdown("##### Ajouter un modèle .docx")
        st.caption("Uploadez vos propres bilans, rapports, notes ou comptes-rendus. L'IA s'en inspirera automatiquement.")

        fichiers_upload = st.file_uploader(
            "Glissez vos fichiers .docx ici",
            type=["docx"],
            accept_multiple_files=True,
            key="upload_modeles"
        )

        if fichiers_upload:
            os.makedirs(DOSSIER_MODELES, exist_ok=True)
            for f in fichiers_upload:
                chemin_dest = os.path.join(DOSSIER_MODELES, f.name)
                with open(chemin_dest, "wb") as out:
                    out.write(f.read())
            st.success(f"✅ {len(fichiers_upload)} modèle(s) ajouté(s) dans inspinote/")
            st.rerun()

        # Liste des modèles existants
        modeles = lister_modeles()
        if modeles:
            st.markdown(f"##### Modèles disponibles ({len(modeles)})")
            for m in modeles:
                col_nom, col_type, col_apercu, col_sup = st.columns([3, 2, 2, 1])
                with col_nom:
                    st.markdown(f"📄 **{m['nom']}**")
                with col_type:
                    st.caption(f"Type : {m['type']}")
                with col_apercu:
                    with st.popover("👁️ Aperçu"):
                        st.text(m["contenu"][:600] + ("..." if len(m["contenu"]) > 600 else ""))
                with col_sup:
                    if st.button("🗑️", key=f"del_modele_{m['fichier']}", use_container_width=True,
                                 help=f"Supprimer {m['nom']}"):
                        try:
                            os.remove(m["chemin"])
                            st.success(f"Modèle {m['nom']} supprimé.")
                            st.rerun()
                        except Exception as e:
                            st.error(f"Erreur : {e}")
        else:
            st.info("Aucun modèle dans inspinote/. Uploadez vos premiers fichiers ci-dessus.")


# ─────────────────────────────────────────────
# MODULE PRINCIPAL
# ─────────────────────────────────────────────
def module_notes(current_user: str):
    inject_css()
    st.markdown("### 📓 Documents d'Entreprise")

    # Gestionnaire de modèles inspinote
    gestionnaire_modeles()

    # Formulaire IA toujours visible en haut
    formulaire_generation_ia(current_user)

    st.markdown("---")

    # Liste des notes sauvegardées
    col_titre, col_refresh = st.columns([5, 1])
    with col_titre:
        st.markdown("#### 📁 Documents sauvegardés")
    with col_refresh:
        if st.button("🔄", use_container_width=True, key="btn_refresh_notes"):
            st.session_state.pop("note_active", None)
            st.rerun()

    notes = get_notes(current_user)

    if not notes:
        st.info("Aucun document sauvegardé. Utilisez l'IA ci-dessus pour en créer un !")
        return

    # Stats
    nb_total = len(notes)
    nb_tab   = sum(1 for n in notes if n.get("tableau") and n["tableau"] not in ["[]",""])
    cats     = {}
    for n in notes:
        c = n.get("categorie","Général")
        cats[c] = cats.get(c,0)+1
    top_cat = max(cats, key=cats.get) if cats else "—"

    c1,c2,c3,c4 = st.columns(4)
    for col,(num,label,color) in zip([c1,c2,c3,c4],[
        (nb_total, "Documents",            "#1E3A8A"),
        (nb_tab,   "Avec tableau",          "#059669"),
        (len(cats),"Catégories",            "#7C3AED"),
        (top_cat,  "Catégorie principale",  "#F59E0B"),
    ]):
        col.markdown(f"""
        <div class="stat-note" style="border-bottom-color:{color}">
            <div class="stat-note-num" style="color:{color}">{num}</div>
            <div class="stat-note-label">{label}</div>
        </div>""", unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)

    col_f1, col_f2 = st.columns([2, 3])
    with col_f1:
        filtre_cat = st.selectbox("Catégorie", ["Toutes"] + CATEGORIES, key="notes_filtre_cat")
    with col_f2:
        recherche  = st.text_input("🔍 Rechercher", placeholder="Mot-clé...", key="notes_recherche")

    notes_filtrees = notes
    if filtre_cat != "Toutes":
        notes_filtrees = [n for n in notes_filtrees if n.get("categorie") == filtre_cat]
    if recherche:
        rech = recherche.lower()
        notes_filtrees = [n for n in notes_filtrees
                          if rech in n.get("titre","").lower() or rech in n.get("contenu","").lower()]

    if not notes_filtrees:
        st.warning("Aucun document ne correspond.")
        return

    note_active_id = st.session_state.get("note_active")
    cols_grille    = st.columns(2)

    for i, note in enumerate(notes_filtrees):
        with cols_grille[i % 2]:
            ouvrir = carte_note(note)
            if ouvrir:
                st.session_state["note_active"] = note["id"]
                st.rerun()
            if note_active_id == note["id"]:
                editeur_note(note, current_user)
