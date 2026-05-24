"""
inspinote.py — Gestionnaire de modèles de documents pour TradAIx
Les modèles .docx sont stockés dans le dossier inspinote/
L'IA lit leur contenu et s'en inspire pour rédiger les documents.
"""

import os
import re
from docx import Document

# ─────────────────────────────────────────────
# CHEMIN DU DOSSIER MODÈLES
# ─────────────────────────────────────────────
DOSSIER_MODELES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "inspinote")

# Correspondance type de document → nom(s) de fichier modèle
MODELES_PAR_TYPE = {
    "📋 Compte-rendu":  ["compte_rendu_reunion.docx"],
    "📊 Rapport":       ["rapport_activite.docx"],
    "📝 Note":          ["note_interne.docx"],
    "📈 Bilan":         ["bilan_financier.docx"],
    "📌 Mémo":          ["memo_rh.docx"],
    "📁 Synthèse":      ["rapport_activite.docx"],
}


# ─────────────────────────────────────────────
# LECTURE D'UN MODÈLE DOCX
# ─────────────────────────────────────────────
def lire_modele_docx(chemin: str) -> str:
    """Extrait le texte complet d'un fichier .docx modèle."""
    try:
        doc  = Document(chemin)
        texte = []
        for para in doc.paragraphs:
            if para.text.strip():
                texte.append(para.text.strip())
        # Lire aussi les tableaux intégrés dans le modèle
        for table in doc.tables:
            for row in table.rows:
                ligne = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if ligne:
                    texte.append(ligne)
        return "\n".join(texte)
    except Exception as e:
        return f"[Erreur lecture modèle : {e}]"


# ─────────────────────────────────────────────
# LISTE DES MODÈLES DISPONIBLES
# ─────────────────────────────────────────────
def lister_modeles() -> list[dict]:
    """Retourne la liste de tous les modèles disponibles dans inspinote/."""
    modeles = []
    if not os.path.exists(DOSSIER_MODELES):
        os.makedirs(DOSSIER_MODELES)
        return modeles

    for fichier in sorted(os.listdir(DOSSIER_MODELES)):
        if fichier.endswith(".docx"):
            chemin  = os.path.join(DOSSIER_MODELES, fichier)
            contenu = lire_modele_docx(chemin)
            # Détecter le type depuis la première ligne [MODELE:xxx]
            type_detecte = "Général"
            match = re.search(r"\[MODELE:(.+?)\]", contenu)
            if match:
                type_detecte = match.group(1).strip()
            # Nom lisible
            nom = fichier.replace("_", " ").replace(".docx", "").title()
            modeles.append({
                "fichier":  fichier,
                "chemin":   chemin,
                "nom":      nom,
                "type":     type_detecte,
                "contenu":  contenu,
                "taille":   len(contenu),
            })
    return modeles


# ─────────────────────────────────────────────
# SUGGESTION AUTOMATIQUE DE MODÈLE
# ─────────────────────────────────────────────
def suggerer_modele(type_doc: str, sujet: str) -> dict | None:
    """Suggère automatiquement le modèle le plus adapté au type de document."""
    modeles      = lister_modeles()
    fichiers_cib = MODELES_PAR_TYPE.get(type_doc, [])

    # 1. Chercher par correspondance exacte de type
    for m in modeles:
        if m["fichier"] in fichiers_cib:
            return m

    # 2. Chercher par type détecté dans le contenu
    type_label = type_doc.split(" ", 1)[-1].lower() if " " in type_doc else type_doc.lower()
    for m in modeles:
        if type_label in m["type"].lower() or type_label in m["nom"].lower():
            return m

    # 3. Chercher par mots-clés du sujet dans le nom du modèle
    mots_sujet = sujet.lower().split()
    for m in modeles:
        for mot in mots_sujet:
            if mot in m["nom"].lower() or mot in m["contenu"].lower()[:200]:
                return m

    # 4. Retourner le premier modèle disponible par défaut
    return modeles[0] if modeles else None


# ─────────────────────────────────────────────
# CONSTRUIRE LE PROMPT D'INSPIRATION
# ─────────────────────────────────────────────
def construire_contexte_modele(modele: dict) -> str:
    """Formate le contenu du modèle pour l'injecter dans le prompt IA."""
    if not modele:
        return ""
    return f"""
--- MODÈLE DE RÉFÉRENCE ({modele['nom']}) ---
Inspire-toi de la structure, du ton et des sections de ce modèle pour rédiger le document.
Adapte le contenu au sujet demandé tout en respectant le format professionnel du modèle.

{modele['contenu'][:3000]}
--- FIN DU MODÈLE ---
"""
